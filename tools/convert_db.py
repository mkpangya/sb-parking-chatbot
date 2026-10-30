"""DB형 엑셀(v3) → 챗봇 데이터(faq.json) 변환 스크립트.

사용법:
    python convert_db.py <DB엑셀.xlsx> <출력경로/faq.json>

동작:
1) _dictionary 시트를 읽어 각 테이블의 필드, 필수 여부, 자료형, 허용 값, 연결 관계를 파악한다.
2) 테이블 시트의 1행이 정의와 같은지 확인하고, 모든 행을 검사한다.
   - 필수값 누락, 기본키 중복, 허용되지 않는 값, 숫자·날짜 형식 오류, 없는 번호 연결
3) 오류가 하나라도 있으면 faq.json을 만들지 않고 목록을 보여준다.
4) 검사를 통과하면 주차장 테이블로부터 주차장별 FAQ 답변을 만들고 faq.json을 저장한다.
"""
import json
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

REQUIRED_SHEETS = ["center", "faq", "faq_synonym", "parking_lot", "discount_type",
                   "discount_note", "contact", "change_log", "menu", "menu_faq"]


# ───────── 읽기·검사 ─────────

def load_dictionary(wb):
    ws = wb["_dictionary"]
    spec = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue
        table, table_label, fname, flabel, ftype, req, key, allowed, refs, desc = (list(row) + [None] * 10)[:10]
        t = spec.setdefault(table, {"label": table_label, "fields": []})
        t["fields"].append({
            "name": fname, "label": flabel, "type": ftype,
            "required": req == "Y", "key": key == "Y",
            "allowed": [x.strip() for x in str(allowed).split(",")] if allowed else [],
            "refs": [x.strip() for x in str(refs).split("|")] if refs else [],
        })
    return spec


def to_date(v):
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip().rstrip(".").replace(".", "-")
    return date.fromisoformat(s)


def read_table(wb, table, fields, errors):
    """시트를 읽어 dict 목록으로 돌려준다. 검사 오류는 errors에 쌓는다."""
    if table not in wb.sheetnames:
        errors.append(f"[{table}] 시트가 없습니다.")
        return []
    ws = wb[table]
    header = [c.value for c in ws[1]]
    expected = [f["name"] for f in fields]
    if header[:len(expected)] != expected:
        errors.append(f"[{table}] 1행 필드명이 정의와 다릅니다. 기대: {expected}")
        return []

    rows = []
    for r_no, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(v not in (None, "") for v in values[:len(expected)]):
            continue
        rec = {}
        for i, f in enumerate(fields):
            v = values[i] if i < len(values) else None
            where = f"[{table}] {r_no}행 {f['label']}({f['name']})"
            if v in (None, ""):
                if f["required"]:
                    errors.append(f"{where}: 필수 값이 비어 있습니다.")
                rec[f["name"]] = None
                continue
            t = f["type"]
            try:
                if t == "int":
                    rec[f["name"]] = int(v)
                elif t == "date":
                    rec[f["name"]] = to_date(v)
                elif t in ("enum", "yn"):
                    s = str(v).strip()
                    if f["allowed"] and s not in f["allowed"]:
                        errors.append(f"{where}: '{s}'은(는) 허용되지 않는 값입니다. ({', '.join(f['allowed'])})")
                    rec[f["name"]] = s
                else:
                    rec[f["name"]] = str(v).strip()
            except (ValueError, TypeError):
                errors.append(f"{where}: 형식이 맞지 않습니다. ({v})")
                rec[f["name"]] = None
        rows.append((r_no, rec))
    return rows


def check_keys(table, fields, rows, errors):
    keys = [f["name"] for f in fields if f["key"]]
    seen = {}
    for r_no, rec in rows:
        k = tuple(rec.get(x) for x in keys)
        if k in seen:
            errors.append(f"[{table}] {r_no}행: 기본키 {keys}={k} 가 {seen[k]}행과 중복됩니다.")
        else:
            seen[k] = r_no


# ───────── 챗봇 답변 만들기 ─────────

def fmt_date(d):
    return d.strftime("%Y.%m.%d") if d else "-"


def num_or_none(v, unit="면"):
    return "없음" if v is None else f"{v}{unit}"


def lot_question(name):
    base = name if name.endswith("주차장") else f"{name} 주차장"
    return f"{base} 정보는 어떻게 되나요?"


def lot_answer(r):
    if r["ev_chargers"] is None:
        chg = "없음"
    else:
        chg = f"{r['ev_chargers']}기"
        if r.get("ev_charger_note"):
            chg += f" ({r['ev_charger_note']})"

    fee = r.get("weekday_fee") or "요금 자료 없음 (주차사업팀 문의)"
    hours = r.get("operating_hours") or "자료에 별도 기재 없음 (주차사업팀 문의)"

    if r["monthly_status"] == "조건부":
        monthly = "자료 확인 중 (주차사업팀 문의)"
    elif r.get("monthly_raw"):
        monthly = r["monthly_raw"]
    else:
        monthly = "자료 없음 (주차사업팀 문의)"

    return (
        f"[{r['name']}]\n"
        f"· 구분: {r['lot_type']}\n"
        f"· 주소: {r['address']}\n"
        f"· 면수: {r['total_spaces']}면\n"
        f"· 전기차 면수: {num_or_none(r['ev_spaces'])}\n"
        f"· 전기차 충전기 수: {chg}\n"
        f"· 장애인 면수: {num_or_none(r['disabled_spaces'])}\n"
        f"· 가족 차량 면수: {num_or_none(r['family_spaces'])}\n"
        f"· 운영개시일: {fmt_date(r['opened_on'])}\n"
        f"· 요금: {fee}\n"
        f"· 운영시간: {hours}\n"
        f"· 월정기(전일/주간/야간): {monthly}"
    )


# ───────── 메인 ─────────

def convert(src: Path, out: Path) -> int:
    wb = load_workbook(src, data_only=True)
    errors, warnings = [], []

    if "_dictionary" not in wb.sheetnames:
        print("[오류] _dictionary 시트가 없습니다. v3 엑셀인지 확인하세요.")
        return 1
    spec = load_dictionary(wb)

    data = {}
    for table in REQUIRED_SHEETS:
        if table not in spec:
            errors.append(f"[_dictionary] '{table}' 정의가 없습니다.")
            continue
        fields = spec[table]["fields"]
        rows = read_table(wb, table, fields, errors)
        check_keys(table, fields, rows, errors)
        data[table] = rows

    # 연결(참조) 검사: _dictionary의 references 정의를 그대로 따른다
    faq_ids = {rec["faq_id"] for _, rec in data.get("faq", [])}
    lot_rows = [rec for _, rec in data.get("parking_lot", [])]
    lot_faq_ids = {r["faq_id"] for r in lot_rows}
    overlap = faq_ids & lot_faq_ids
    for x in sorted(overlap):
        errors.append(f"[parking_lot] 안내 FAQ 번호 '{x}'이 faq 시트와 겹칩니다.")
    for table in REQUIRED_SHEETS:
        if table not in spec:
            continue
        for f in spec[table]["fields"]:
            if not f["refs"]:
                continue
            targets = set()
            for ref in f["refs"]:
                t_name, t_field = ref.split(".")
                targets |= {rec[t_field] for _, rec in data.get(t_name, []) if rec.get(t_field)}
            for r_no, rec in data.get(table, []):
                v = rec.get(f["name"])
                if v is not None and v not in targets:
                    errors.append(f"[{table}] {r_no}행 {f['label']}({f['name']}): '{v}'이(가) 연결 대상 "
                                  f"({', '.join(f['refs'])})에 없습니다.")

    if errors:
        print("[오류] 아래 항목을 고친 뒤 다시 실행하세요. faq.json은 만들지 않았습니다.")
        for e in errors:
            print("  -", e)
        return 1

    # 게시용 데이터 만들기
    faqs = []
    for _, rec in data["faq"]:
        if rec["publish_yn"] != "Y":
            continue
        faqs.append({
            "id": rec["faq_id"], "center": rec["center"], "category": rec["category"], "question": rec["question"],
            "answer": rec["answer"], "source": rec["source"], "as_of": rec["as_of"],
            "status": rec["status"],
        })

    def map_places(lot_type=None):
        """지도 버튼용 주차장 목록 (게시된 주차장만, lot_type이 있으면 해당 구분만)"""
        return [{"name": r["name"], "address": r["address"], "lot_type": r["lot_type"]}
                for r in lot_rows
                if r["publish_yn"] == "Y" and (lot_type is None or r["lot_type"] == lot_type)]

    for r in lot_rows:
        if r["publish_yn"] != "Y":
            continue
        faqs.append({
            "id": r["faq_id"], "center": "공영주차센터", "category": "주차장 현황",
            "question": lot_question(r["name"]),
            "answer": lot_answer(r), "source": "공단 현황자료(2026.09)", "as_of": "2026.09",
            "status": r["faq_status"],
            "places": [{"name": r["name"], "address": r["address"], "lot_type": r["lot_type"]}],
        })

    # 주소를 나열하는 FAQ에는 카카오지도 버튼용 주차장 목록을 붙인다 (FAQ 번호 → 구분 필터, None은 전체)
    MAP_FAQ_FILTER = {"F-051": None, "F-031": "노상", "F-032": "노외", "F-033": "부설"}
    for f in faqs:
        if f["id"] in MAP_FAQ_FILTER:
            f["places"] = map_places(MAP_FAQ_FILTER[f["id"]])

    synonyms = {}
    # 동의어는 게시 여부와 관계없이 모두 저장한다. 나중에 게시(Y)로 바꾸면 바로 검색되도록 하기 위함이다.
    # 위젯은 게시된 FAQ만 화면에 보여준다.
    for _, rec in data["faq_synonym"]:
        synonyms.setdefault(rec["keyword"], [])
        if rec["faq_id"] not in synonyms[rec["keyword"]]:
            synonyms[rec["keyword"]].append(rec["faq_id"])

    faqs.sort(key=lambda f: f["id"])  # 항상 같은 순서로 저장

    # 메인 메뉴 (게시된 메뉴만, 연결된 FAQ는 게시된 것만)
    published_faq = {f["id"]: f for f in faqs}
    menus = []
    for _, m in sorted(data["menu"], key=lambda x: x[1]["sort_order"]):
        if m["publish_yn"] != "Y":
            continue
        items = []
        if m["action"] == "faq_list":
            links = sorted([rec for _, rec in data["menu_faq"] if rec["menu_id"] == m["menu_id"]],
                           key=lambda x: x["sort_order"])
            for link in links:
                f = published_faq.get(link["faq_id"])
                if f is None:
                    warnings.append(f"메뉴 '{m['label']}': {link['faq_id']}는 게시되지 않아 목록에서 제외됩니다.")
                    continue
                items.append({"id": f["id"], "question": f["question"]})
        menus.append({"id": m["menu_id"], "label": m["label"], "sub": m["sub_label"] or "",
                      "icon": m["icon"], "action": m["action"], "items": items})

    centers = [{"name": r["center_name"], "phone": r["phone"] or "", "hours": r["hours"] or "",
                "scope": r["scope"]} for _, r in data["center"]]
    by_center = {}
    for f in faqs:
        by_center[f["center"]] = by_center.get(f["center"], 0) + 1
    print("  센터별 게시 FAQ: " + ", ".join(f"{k} {by_center.get(k, 0)}개" for k in
                                       ["공통", "거주자주차센터", "공영주차센터"]))

    contacts = [{"name": r["name"], "phone": r["phone"], "hours": r["hours"], "purpose": r["purpose"]}
                for _, r in data["contact"] if r["publish_yn"] == "Y"]

    discounts = []
    for _, r in sorted(data["discount_type"], key=lambda x: x[1]["sort_order"]):
        if r["publish_yn"] == "Y":
            discounts.append({"type": r["name"], "rate": r["rate_text"] or "", "proof": r["proof"]})
    for _, r in sorted(data["discount_note"], key=lambda x: x[1]["sort_order"]):
        if r["publish_yn"] == "Y":
            discounts.append({"type": r["note_group"], "rate": "", "proof": r["note_text"]})

    # 합계 점검 (챗봇에는 표시하지 않음, 검토용 출력)
    tot = {
        "면수": sum(r["total_spaces"] or 0 for r in lot_rows),
        "전기차 면수": sum(r["ev_spaces"] or 0 for r in lot_rows),
        "충전기(구청부설 포함)": sum(r["ev_chargers"] or 0 for r in lot_rows),
        "장애인 면수": sum(r["disabled_spaces"] or 0 for r in lot_rows),
        "가족 차량 면수": sum(r["family_spaces"] or 0 for r in lot_rows),
    }

    for msg in warnings:
        print("[경고]", msg)
    print("검사 통과")
    print(f"  테이블 행 수: faq {len(data['faq'])}, 동의어 {len(data['faq_synonym'])}, "
          f"주차장 {len(lot_rows)}, 감면 {len(data['discount_type'])}, 감면 유의 {len(data['discount_note'])}, "
          f"문의처 {len(data['contact'])}, 변경이력 {len(data['change_log'])}")
    print(f"  챗봇 게시 FAQ: {len(faqs)}개, 동의어: {len(synonyms)}개, 감면 표 행: {len(discounts)}개, 메뉴: {len(menus)}개")
    print(f"  주차장 합계 점검: " + ", ".join(f"{k} {v}" for k, v in tot.items()))

    result = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "faq_count": len(faqs),
        "faqs": faqs,
        "synonyms": synonyms,
        "contacts": contacts,
        "discounts": discounts,
        "menus": menus,
        "centers": centers,
    }
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"저장 완료: {out}")
    return 0


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("사용법: python convert_db.py <DB엑셀.xlsx> <출력경로/faq.json>")
        sys.exit(1)
    sys.exit(convert(Path(sys.argv[1]), Path(sys.argv[2])))
